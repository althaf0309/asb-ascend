"""Convert the supplied logistics workbook into deterministic SEO page data."""
from pathlib import Path
import json
import re
import sys
import unicodedata
import openpyxl

SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])

def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip(" -")

def slugify(value):
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")

def intent_for(title):
    value = title.lower()
    if value.endswith("?") or value.startswith(("what ", "which ", "where ", "how ", "is ", "can ", "does ", "who ")):
        return "question"
    for needle, intent in [
        ("fee", "fees"), ("cost", "fees"), ("admission", "admission"),
        ("placement", "placement"), ("internship", "internship"),
        ("after 12", "after-school"), ("plus two", "after-school"),
        ("after degree", "graduate"), ("career", "career"), ("job", "career"),
        ("salary", "career"), ("online", "online"), ("near me", "near-me"),
        ("certif", "certification"), ("institute", "institute"),
        ("best", "comparison"), ("duration", "duration"),
    ]:
        if needle in value:
            return intent
    return "course"

book = openpyxl.load_workbook(SOURCE, data_only=True)
sheet = book.active
families = [(1, "logistics", "Logistics and Supply Chain Management"), (2, "warehouse", "Warehouse and Inventory Management")]
pages = []

for column, family, label in families:
    seen_titles, seen_slugs = set(), set()
    for row in range(1, sheet.max_row + 1):
        title = clean(sheet.cell(row, column).value)
        if not title or title.lower().startswith("http") or "url based" in title.lower():
            continue
        identity = title.casefold()
        if identity in seen_titles:
            continue
        seen_titles.add(identity)
        slug = slugify(title)
        if not slug:
            continue
        original = slug
        suffix = 2
        while slug in seen_slugs:
            slug = f"{original}-{suffix}"
            suffix += 1
        seen_slugs.add(slug)
        location = ""
        match = re.search(r"\bin\s+([A-Za-z][A-Za-z .'-]+?)(?:\?|$)", title, re.I)
        if match:
            candidate = clean(match.group(1)).title()
            course_words = ("logistics", "warehouse", "supply chain", "inventory", "management", "course", "diploma", "certification", "training")
            if not any(word in candidate.lower() for word in course_words):
                location = candidate
        pages.append({
            "family": family,
            "slug": slug,
            "title": title[0].upper() + title[1:],
            "label": label,
            "intent": intent_for(title),
            "location": location,
            "question": intent_for(title) == "question",
        })

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(pages, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Generated {len(pages)} pages: " + ", ".join(f"{f}={sum(p['family']==f for p in pages)}" for _,f,_ in families))
